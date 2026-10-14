"""Audit author label dependencies without inventing states or reading targets."""
import argparse
import json
from pathlib import Path
import subprocess
from urllib.parse import quote

from openpyxl import load_workbook
from veloroute.artifacts import Run, save_csv, save_json


def gh(path, raw=False):
    command = ['gh', 'api', path]
    if raw:
        command += ['-H', 'Accept: application/vnd.github.raw+json']
    value = subprocess.check_output(command, text=True, timeout=45)
    return value if raw else json.loads(value)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--output', required=True)
    args = ap.parse_args()
    root = Path('/data/yuchang/veloroute_ucheck_20260915')
    repo = 'Yeqing95/MeRLin'
    commit = gh(f'repos/{repo}/commits/HEAD')['sha']
    tree = gh(f'repos/{repo}/git/trees/{commit}?recursive=1')
    if tree.get('truncated'):
        raise ValueError('Incomplete author repository inventory')
    files = [v for v in tree['tree'] if v['type'] == 'blob']
    script_name = '3.ClonoCluster.R'
    author_script = gh(f'repos/{repo}/contents/{quote(script_name)}?ref={commit}', raw=True)
    required = [f'group{i}_bc.txt' for i in range(1, 6)]
    if any(name not in author_script for name in required):
        raise ValueError('Author label procedure changed: inspect before interpreting')
    local_roots = [root, Path('/data/yuchang/caixujia/merlin_v1')]
    hits = {name: [str(p) for base in local_roots if base.exists() for p in base.rglob(name)] for name in required}
    repo_names = {Path(v['path']).name for v in files}
    missing = [name for name in required if name not in repo_names and not hits[name]]
    config = dict(author_repo=repo, author_commit=commit, searched_local_roots=[str(p) for p in local_roots],
                  scope='dependency_inventory_not_fate_label_generation',
                  git_commit=subprocess.check_output(['git', 'rev-parse', 'HEAD'], text=True).strip())
    with Run(args.output, stage='W3_6b_author_label_inventory', kind='engineering', config=config,
             inputs=[Path(__file__), root / 'merlin_programs.json', root / 'merlin_suppl.xlsx']) as run:
        save_json(run.directory / 'author_tree.json', tree)
        (run.directory / 'author_ClonoCluster.R').write_text(author_script)
        save_json(run.directory / 'local_label_matches.json', hits)
        wb = load_workbook(root / 'merlin_suppl.xlsx', read_only=True, data_only=True)
        sheets = [dict(sheet=ws.title, rows=ws.max_row, columns=ws.max_column, title=str(ws.cell(1, 1).value)) for ws in wb]
        wb.close()
        save_csv(run.directory / 'supplementary_inventory.csv', sheets)
        result = dict(status='AUTHOR_GROUP_FILES_MISSING_IN_CHECKED_LOCATIONS' if missing else 'LABEL_FILES_REQUIRE_ALIGNMENT_AUDIT',
            author_repo=repo, author_commit=commit, required_group_files=required, missing=missing,
            labels_generated=False, future_expression_read=False, github_files=len(files),
            interpretation='Author bc_group is assigned by clone Barcode membership, not argmax of five signatures. '
                'With this mapping every clone has one group by construction; it is not independently measured within-clone fate composition.',
            remaining_search='GEO provides GSM9044665_D21_Clonocluster.rds (500555118 bytes), '
                'an alternative source for bc_group metadata; acquisition/alignment is separate. No author contact sent')
        save_json(run.directory / 'summary.json', result)
        (run.directory / 'RESULTS.md').write_text('# MeRLin author label dependency audit\n\n' + json.dumps(result, indent=2) + '\n')
        print(json.dumps(result), flush=True)


if __name__ == '__main__':
    main()
