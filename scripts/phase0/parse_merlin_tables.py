"""Parse MeRLin supplementary tables into: barcode extraction config + program signatures."""
import json
from pathlib import Path

import openpyxl

ROOT = Path('/data/yuchang/veloroute_ucheck_20260915')
wb = openpyxl.load_workbook(ROOT/'merlin_suppl.xlsx', read_only=True, data_only=True)


def column_values(ws, column, start_row, max_blank=30):
    values, blanks = [], 0
    for row in ws.iter_rows(min_row=start_row, max_col=column, values_only=True):
        value = row[column-1] if len(row) >= column else None
        if value is None or str(value).strip() == '':
            blanks += 1
            if blanks >= max_blank:
                break
            continue
        blanks = 0
        values.append(str(value).strip())
    return values


def table1_barcode_config():
    ws = wb['Supplementary Table 1']
    rows = {}
    for row in ws.iter_rows(min_row=1, max_row=80, values_only=True):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if len(cells) >= 2:
            rows[cells[0]] = cells[1]
    config = dict(
        flank_3prime_vector_backbone='CAGATCTTAGCCACTTTTTAAAAGAAAAGGGGG',
        source='paper Methods: extractor scans 3prime vector backbone to locate barcode start; Table S1',
        barcode_first_bases=15,
        insert_1='BC_Ins_1_RC='+rows.get('BC_Ins_1_RC', ''),
        insert_2='BC_Ins_2='+rows.get('BC_Ins_2', ''),
        qc=['remove_empty_vectors', 'remove_short_barcodes', 'remove_incorrect_patterns',
            'remove_cells_with_multiple_barcodes'],
        notes='Barcode immediately follows the constant flank in read orientation; '
              'pattern letters in inserts are semi-random (W=A/T, S=C/G, N=any).')
    return config


SIGNATURE_NAMES = ['group1_stress_like', 'groups23_neural_crest_like', 'group2_lipid_metabolism',
                   'group3_pi3k_signaling', 'group4_ecm_remodeling']


def table9_signatures():
    ws = wb['Supplementary Table 9']
    signatures = {}
    for i, name in enumerate(SIGNATURE_NAMES, start=1):
        genes = column_values(ws, i, start_row=7)
        signatures[name] = genes
    return signatures


def table8_degs():
    ws = wb['Supplementary Table 8']
    headers = {}
    for row in ws.iter_rows(min_row=6, max_row=6, values_only=True):
        for i, value in enumerate(row, start=1):
            if value is not None and 'group' in str(value).lower():
                headers[i] = str(value).strip()
    sections = {}
    for row in ws.iter_rows(min_row=4, max_row=4, values_only=True):
        for i, value in enumerate(row, start=1):
            if value is not None and str(value).strip():
                sections[str(value).strip()] = i
    out = {}
    for column, header in headers.items():
        section = 'unknown'
        for name, start in sorted(sections.items(), key=lambda kv: kv[1]):
            if column >= start:
                section = name
        key = f"{section}::{header}"
        out[key] = column_values(ws, column, start_row=7)
    return out


def main():
    barcode = table1_barcode_config()
    (ROOT/'merlin_barcode_config.json').write_text(json.dumps(barcode, indent=2))
    signatures = table9_signatures()
    degs = table8_degs()
    programs = dict(signatures=signatures, degs=degs,
                    source='Tables S8/S9; DEGs from BRAFi/MEKi-treated WM4237, FC>=2, cell%>=50%, FDR<0.05')
    (ROOT/'merlin_programs.json').write_text(json.dumps(programs, indent=2))
    print('barcode config:', json.dumps({k: v for k, v in barcode.items() if k != 'notes'}, indent=1)[:600])
    for name, genes in signatures.items():
        print(f'{name}: {len(genes)} genes, first: {genes[:5]}')
    for name, genes in list(degs.items())[:8]:
        print(f'{name}: {len(genes)} genes')
    print('total deg columns:', len(degs))


if __name__ == '__main__':
    main()
